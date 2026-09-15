"""Regenerate the CSV caches the readtable shim reads, from capInfo.xlsx.

Values are written with repr() so the doubles round-trip exactly.
Run from the ROAST root:  python3 octave-compat/scripts/export_capinfo.py
"""
import os
import openpyxl

out_dir = os.path.join(os.path.dirname(__file__), "..", "shims")
wb = openpyxl.load_workbook("capInfo.xlsx", data_only=True)
for sheet in wb.sheetnames:
    path = os.path.join(out_dir, "capInfo__%s.csv" % sheet.replace("/", "_"))
    with open(path, "w") as fh:
        for row in wb[sheet].iter_rows(values_only=True):
            if row[0] is None:
                continue
            fh.write("%s,%r,%r,%r\n" % (str(row[0]), float(row[1]),
                                        float(row[2]), float(row[3])))
    print("wrote", path)
